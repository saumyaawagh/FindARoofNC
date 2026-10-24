"""Step 00: check that every data source is present and has the columns we expect.

Prints one line per source:
  ✅ present   ⚠️ optional and missing   ⏳ downloaded by a later step   ❌ required and missing
Exits with an error if anything is ❌, so run_all.py stops before building on bad inputs.
How to get each manual file: pipeline/manual/README.md
"""
import re
import sys
import warnings

import openpyxl
import pandas as pd

import config

# openpyxl warns about Excel styling it can't read; that doesn't affect the data.
warnings.filterwarnings("ignore", module="openpyxl")


def header_of(path, source):
    """Return the list of column names in a CSV or Excel file without loading all of it."""
    if path.suffix == ".csv":
        return list(pd.read_csv(path, nrows=0, encoding="utf-8-sig").columns)
    wb = openpyxl.load_workbook(path, read_only=True)
    try:
        ws = wb[source["sheet"]]
    except KeyError:
        raise ValueError(f"no sheet named {source['sheet']!r} (sheets: {wb.sheetnames})")
    skip = source.get("header_row", 0)
    for i, row in enumerate(ws.iter_rows(max_row=skip + 1, values_only=True)):
        if i == skip:
            return [str(c) for c in row if c is not None]
    return []


def check_civil_fiscal_year(path):
    """The file name says which fiscal year it should hold; make sure the dates inside agree.

    (In Phase 1 the two downloads arrived with swapped names, so this is a real risk.)
    """
    m = re.search(r"fy(\d{2})-(\d{2})", path.name)
    wb = openpyxl.load_workbook(path, read_only=True)
    first = next(wb[config.CIVIL_SHEET].iter_rows(min_row=2, max_row=2, values_only=True))
    start, end = first[0], first[1]
    expected = (2000 + int(m.group(1)), 2000 + int(m.group(2)))
    if (start.year, end.year) != expected:
        raise ValueError(
            f"file name says FY {expected[0]}-{expected[1]} but it covers "
            f"{start:%Y-%m-%d} to {end:%Y-%m-%d}; rename the file to match its contents"
        )
    return f"FY {start.year}-{str(end.year)[2:]} ({start:%b %Y}–{end:%b %Y})"


def check_foreclosure_title(path):
    """Report how far the file runs (its title says 'THROUGH <date>')."""
    wb = openpyxl.load_workbook(path, read_only=True)
    title = next(wb[config.FORECLOSURE_SHEET].iter_rows(max_row=1, values_only=True))[0]
    m = re.search(r"THROUGH (\w+ \d+, \d{4})", str(title))
    return f"through {m.group(1)}" if m else "date range not stated in title"


EXTRA_CHECKS = {
    "nc_evictions": check_civil_fiscal_year,
    "nc_foreclosures": check_foreclosure_title,
}


def check_source(source, present_ids):
    """Return (symbol, note) for one manifest entry."""
    missing = [p for p in source["files"] if not p.exists()]

    if source["kind"] == "scripted":
        if missing:
            return "⏳", f"will be downloaded by {source['fetched_by']}"
        return "✅", f"downloaded by {source['fetched_by']}"

    required = not source["optional"] or source.get("required_if") in present_ids
    if missing:
        names = ", ".join(p.name for p in missing)
        if required:
            why = f" (needed because {source['required_if']} is present)" if source.get("required_if") else ""
            return "❌", f"missing {names}{why}"
        return "⚠️", f"optional, missing {names}"

    notes = []
    for path in source["files"]:
        expected = source["columns"].get(path.name, [])
        if expected:
            try:
                cols = header_of(path, source)
            except ValueError as err:
                return "❌", f"{path.name}: {err}"
            absent = [c for c in expected if c not in cols]
            if absent:
                return "❌", f"{path.name} is missing columns {absent}"
        if source["id"] in EXTRA_CHECKS:
            try:
                notes.append(f"{path.name}: {EXTRA_CHECKS[source['id']](path)}")
            except ValueError as err:
                return "❌", f"{path.name}: {err}"
    return "✅", "; ".join(notes) or "present"


def main():
    # Which optional sources are here? Crosswalks are only required when their data source is.
    present_ids = {s["id"] for s in config.SOURCES if all(p.exists() for p in s["files"])}

    print(f"{'':2} {'source':<18} {'kind':<9} status")
    print("-" * 78)
    failures = 0
    for source in config.SOURCES:
        symbol, note = check_source(source, present_ids)
        failures += symbol == "❌"
        print(f"{symbol:2} {source['id']:<18} {source['kind']:<9} {note}")

    print("-" * 78)
    if failures:
        print(f"{failures} required source(s) missing or malformed. See pipeline/manual/README.md.")
        sys.exit(1)
    print("All required sources present.")


if __name__ == "__main__":
    main()
